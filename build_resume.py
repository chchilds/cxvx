#!/usr/bin/env python3
"""Generate Chris Childs 2026 resume (v14) as a Word document."""

import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

NAVY = RGBColor(0x1B, 0x36, 0x5D)
BLACK = RGBColor(0x22, 0x22, 0x22)
GRAY = RGBColor(0x44, 0x44, 0x44)
RULE = "1B365D"

OUT = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "Resume - Chris Childs 2026 (v14).docx",
)


def set_run_font(run, name="Calibri", size=11, bold=False, italic=False, color=BLACK):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def set_paragraph_spacing(p, before=0, after=0, line=1.08, exact=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if exact is not None:
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(exact)
    else:
        pf.line_spacing = line
    pf.widow_control = True


def add_bottom_border(paragraph, color=RULE, size="12"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_hyperlink(paragraph, text, url, size=10.5, color="1B365D"):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), "Calibri")
    rFonts.set(qn("w:hAnsi"), "Calibri")
    rPr.append(rFonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rPr.append(sz)
    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), str(int(size * 2)))
    rPr.append(szCs)
    color_el = OxmlElement("w:color")
    color_el.set(qn("w:val"), color)
    rPr.append(color_el)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    new_run.append(rPr)
    text_el = OxmlElement("w:t")
    text_el.set(qn("xml:space"), "preserve")
    text_el.text = text
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def section_heading(doc, text):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=8, after=3, exact=15)
    run = p.add_run(text.upper())
    set_run_font(run, size=12, bold=True, color=NAVY)
    add_bottom_border(p)
    return p


def job_header(doc, title, dates):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=5, after=0, exact=14)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.tab_stops.add_tab_stop(Inches(7.2), WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run(title)
    set_run_font(r, size=11, bold=True, color=BLACK)
    r = p.add_run("\t" + dates)
    set_run_font(r, size=10.5, italic=True, color=GRAY)
    return p


def company_line(doc, company, location):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=0, after=1, exact=13)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(company)
    set_run_font(r, size=10.5, italic=True, color=NAVY)
    r = p.add_run(f"  |  {location}")
    set_run_font(r, size=10.5, italic=False, color=GRAY)
    return p


def bullet(doc, text, keep_next=True):
    p = doc.add_paragraph(style="List Bullet")
    set_paragraph_spacing(p, before=0, after=0, exact=13)
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.15)
    p.paragraph_format.keep_with_next = keep_next
    # Clear default run if any
    if p.runs:
        p.runs[0].text = ""
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=BLACK)
    return p


def body(doc, text, size=10.5, italic=False, before=2, after=4):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=before, after=after, exact=14)
    r = p.add_run(text)
    set_run_font(r, size=size, italic=italic, color=BLACK)
    return p


def edu_line(doc, left, right):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=2, after=0, exact=14)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.tab_stops.add_tab_stop(Inches(7.2), WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run(left)
    set_run_font(r, size=10.5, color=BLACK)
    r = p.add_run("\t" + right)
    set_run_font(r, size=10.5, italic=True, color=GRAY)
    return p


def skill_row(doc, label, items):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=1, after=1, exact=14)
    r = p.add_run(f"{label}:  ")
    set_run_font(r, size=10.5, bold=True, color=NAVY)
    r = p.add_run(items)
    set_run_font(r, size=10.5, color=BLACK)
    return p


def build():
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    # Default style
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BLACK
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    # Bullet style
    bullet_style = doc.styles["List Bullet"]
    bullet_style.font.name = "Calibri"
    bullet_style.font.size = Pt(10.5)
    pf = bullet_style.paragraph_format
    pf.left_indent = Inches(0.25)
    pf.first_line_indent = Inches(-0.15)

    # Footer
    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(fp, before=0, after=0)
    r = fp.add_run("Christopher E. Childs  ·  Resume 2026 (v14)  ·  Page ")
    set_run_font(r, size=8.5, color=GRAY)
    # PAGE field
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_text = OxmlElement("w:t")
    fld_text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run = fp.add_run()
    set_run_font(run, size=8.5, color=GRAY)
    run._r.append(fld_begin)
    r2 = fp.add_run()
    set_run_font(r2, size=8.5, color=GRAY)
    r2._r.append(instr)
    r3 = fp.add_run()
    set_run_font(r3, size=8.5, color=GRAY)
    r3._r.append(fld_sep)
    r3._r.append(fld_text)
    r4 = fp.add_run()
    set_run_font(r4, size=8.5, color=GRAY)
    r4._r.append(fld_end)
    r = fp.add_run(" of ")
    set_run_font(r, size=8.5, color=GRAY)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " NUMPAGES "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_text = OxmlElement("w:t")
    fld_text.text = "2"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run = fp.add_run()
    set_run_font(run, size=8.5, color=GRAY)
    run._r.append(fld_begin)
    r2 = fp.add_run()
    set_run_font(r2, size=8.5, color=GRAY)
    r2._r.append(instr)
    r3 = fp.add_run()
    set_run_font(r3, size=8.5, color=GRAY)
    r3._r.append(fld_sep)
    r3._r.append(fld_text)
    r4 = fp.add_run()
    set_run_font(r4, size=8.5, color=GRAY)
    r4._r.append(fld_end)

    # ---- Header ----
    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(name, before=0, after=1, exact=28)
    r = name.add_run("CHRISTOPHER E. CHILDS")
    set_run_font(r, size=22, bold=True, color=NAVY)

    tag = doc.add_paragraph()
    tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(tag, before=0, after=4, exact=14)
    r = tag.add_run(
        "Premium Services Expert  ·  Strategic Software & Services Architect  ·  Cisco Sales Champion"
    )
    set_run_font(r, size=11, italic=True, color=GRAY)

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(contact, before=0, after=2, exact=13)
    r = contact.add_run("2632 Gross Ave, Wake Forest, North Carolina 27587  ·  (919) 523-8426")
    set_run_font(r, size=10.5, color=BLACK)

    emails = doc.add_paragraph()
    emails.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(emails, before=0, after=2, exact=13)
    add_hyperlink(emails, "Chchilds@cisco.com", "mailto:Chchilds@cisco.com")
    r = emails.add_run("  ·  ")
    set_run_font(r, size=10.5, color=BLACK)
    add_hyperlink(emails, "Chris.E.Childs@gmail.com", "mailto:Chris.E.Childs@gmail.com")
    r = emails.add_run("  ·  ")
    set_run_font(r, size=10.5, color=BLACK)
    add_hyperlink(
        emails,
        "linkedin.com/in/chrischildsraleigh",
        "https://www.linkedin.com/in/chrischildsraleigh",
    )

    # ---- Summary ----
    section_heading(doc, "Professional Summary")
    body(
        doc,
        "Customer-focused Cisco combined seller and solutions engineer with 14+ years at Cisco spanning "
        "premium services, strategic software, Enterprise Agreements, and pre-sales architecture. "
        "Designs consumption models that attach CX Enterprise Agreements (CXEA), Expert Care, Lifecycle Services, "
        "and Advanced Services to software ELAs — and converts that design work into competitive wins across "
        "security, collaboration, data center, wireless, and enterprise networking.",
        before=3,
        after=1,
    )

    # ---- Experience ----
    section_heading(doc, "Experience")

    company_line(doc, "Cisco Systems, Inc.", "RTP, NC  ·  Denver, CO  ·  United States")

    job_header(
        doc,
        "EA Architect & Premium Service Seller (Combined Seller)",
        "August 2023 – Present",
    )
    bullet(
        doc,
        "Combined seller for Cisco Premium Services and Strategic Software: architect CX Enterprise Agreements "
        "(CXEAs), Expert Care, Lifecycle Services (LCS), and Advanced Services (AS) offers as a single customer motion.",
    )
    bullet(
        doc,
        "Design software-plus-services consumption models that attach Premium Services to Enterprise Agreements "
        "and drive recurring-revenue, True Forward, and high-adoption outcomes.",
    )
    bullet(
        doc,
        "Partner with account teams, partners, and customers to position One Cisco architectures across "
        "networking, security, collaboration, and data center software.",
    )
    bullet(
        doc,
        "Recognized as a Cisco Sales Champion and GSX winner for software and services sales; "
        "completed IMPACT Ready IC (Aug 2023) and FY24 Services Sales (Oct 2023).",
        keep_next=False,
    )

    job_header(doc, "Strategic Software Sales Specialist", "October 2019 – Present")
    bullet(
        doc,
        "Lead strategic software sales for Enterprise Agreements, ELA consumption, and subscription / "
        "recurring-revenue offers across Cisco’s software portfolio.",
    )
    bullet(
        doc,
        "Named FY22 Americas Sales Champion, Services & Software — recognized among Cisco’s top "
        "Americas software and services sellers.",
    )
    bullet(
        doc,
        "Translate security, collaboration, networking, and data center software into customer consumption "
        "and adoption plans; completed Cisco Recurring Revenue Metrics (Feb 2022).",
    )
    bullet(
        doc,
        "Carry forward SE-honed competitive displacement skills — SD-WAN, Meraki, security, contact center, "
        "and API/scripting use cases — into software-led deals.",
        keep_next=False,
    )

    job_header(doc, "Systems Engineer, Carolina Select", "January 2015 – October 2019")
    bullet(
        doc,
        "Technical and engineering lead for Cisco Select healthcare, financial, and manufacturing accounts in "
        "North and South Carolina. Competitive wins in security, collaboration, data center, wireless, enterprise, "
        "management, and software, delivered with Advanced Services or partner-led engagements.",
    )
    bullet(
        doc,
        "Drove blended 100% ELA / software consumption spanning UC, ONE (C1), and Security; used CloudCenter, ACI, "
        "and UCS Director in competitive data-center-to-cloud conversations modeled on Cisco-on-Cisco catalog offerings.",
    )
    bullet(
        doc,
        "SE of multiple quarters in Commercial; seven Carolina Select customers; customer-referenced architectures, "
        "EBCs, Global CRO travel, Cisco Live Expo for LED / UPOE with CREE, plus scripting and API use cases.",
    )
    bullet(
        doc,
        "Major wins in SD-WAN, Meraki, Security portfolio, and Contact Center; full post-sales and customer-excellence focus.",
        keep_next=False,
    )

    job_header(doc, "Network Consulting Engineer", "March 2013 – January 2015")
    bullet(
        doc,
        "Designed and implemented WAAS and ACE for production data centers (HLD, LLD, NIP); 90%+ customer-facing time; "
        "six Advanced Services enterprise customers; UCS deploy and on-site knowledge transfers with TAC and BU partners.",
        keep_next=False,
    )

    job_header(doc, "Network Engineer, Remote Management Services", "May 2012 – March 2013")
    bullet(
        doc,
        "NOC incident handling for Cisco premier customers (Remedy, Cisco MAP, TAC): LAN/WLAN, Nexus 7000 NX-OS, "
        "ASR5k 3G, ASR9k IOS XR LTE, backhaul, TACACS+/ACS, IronPort, IPsec/VPN, BGP/OSPF/EIGRP; built handover, "
        "outage-mailer, and NMS-access processes with SLA/SLO reporting.",
        keep_next=False,
    )

    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=6, after=1, exact=14)
    p.paragraph_format.keep_with_next = True
    r = p.add_run("Earlier career")
    set_run_font(r, size=11, bold=True, color=BLACK)

    bullet(
        doc,
        "Realtor, Smart Choice Realty and ERA United Home Realty, Raleigh, NC (2006–2012) — built and maintained "
        "the firm website and email server; client database search, digital photos, PDF contracts, and e-sign storage.",
    )
    bullet(
        doc,
        "Sales Associate / Geek Squad, Best Buy, Cary, NC (2004–2006) — computer sales and home/small-business "
        "networking support; MVP award for Exemplary Service.",
    )
    bullet(
        doc,
        "General Manager, Papa John’s, Cary, NC (2000–2004) — POS/credit-card connectivity, customer database, "
        "web-ordering coupon sync, and productivity / P&L / inventory reporting.",
        keep_next=False,
    )

    # ---- Education ----
    section_heading(doc, "Education")
    edu_line(
        doc,
        "East Carolina University — B.S., Industrial Technology",
        "Completed 2018",
    )
    edu_line(
        doc,
        "Wake Technical Community College — A.A.S., Networking Technology, GPA 3.7",
        "2013",
    )
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=1, after=0, exact=13)
    r = p.add_run(
        "Wake Tech: student leader / peer tutor; Disability Support Services presentation on transitioning for "
        "students with disabilities.  ECU: recruiting for Cisco Advanced Services internships."
    )
    set_run_font(r, size=10, italic=True, color=GRAY)

    # ---- Certifications ----
    section_heading(doc, "Certifications & Training")
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=3, after=1, exact=14)
    r = p.add_run("Current / recent:  ")
    set_run_font(r, size=10.5, bold=True, color=NAVY)
    r = p.add_run(
        "FY24 Services Sales (Oct 2023); IMPACT Ready ICs (Aug 2023); Cisco Recurring Revenue Metrics (Feb 2022)."
    )
    set_run_font(r, size=10.5, color=BLACK)

    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=1, after=1, exact=14)
    r = p.add_run("Previously held:  ")
    set_run_font(r, size=10.5, bold=True, color=NAVY)
    r = p.add_run(
        "CCNP Routing & Switching, CCDP Design, CCNP Data Center, CCNA Voice (valid through 2020); "
        "VCP 6.0 DCV (valid through 2018); CCIE Data Center written passed (lab date Nov 2017)."
    )
    set_run_font(r, size=10.5, color=BLACK)

    # ---- Skills ----
    section_heading(doc, "Skills")
    skill_row(
        doc,
        "Software & services sales",
        "Enterprise Agreements / ELA, CXEA, Premium Services, Expert Care, LCS, Advanced Services, "
        "consumption modeling, recurring revenue, True Forward, One Cisco",
    )
    skill_row(
        doc,
        "Architecture",
        "SD-WAN, Meraki, Security portfolio, Contact Center / UC, ACI, UCS, CloudCenter, wireless, campus / WAN, UPOE",
    )
    skill_row(
        doc,
        "Engineering",
        "Routing & switching (BGP, OSPF, EIGRP, STP/VLANs), NX-OS / IOS / IOS XR, ASA / IronPort, IPsec VPN, QoS, IPv6, packet analysis",
    )
    skill_row(
        doc,
        "Automation & tools",
        "Scripting and Cisco APIs, Wireshark, SNMP, Remedy, documentation, customer-referenced architectures",
    )
    skill_row(
        doc,
        "Personal",
        "Customer-facing, multi-account ownership, solution-oriented, communicative, time management, flexible",
    )

    # Core properties
    core = doc.core_properties
    core.author = "Christopher E. Childs"
    core.last_modified_by = "Christopher E. Childs"
    core.title = "Resume — Christopher E. Childs 2026 (v14)"
    core.subject = "Cisco EA Architect, Premium Service Seller, and Strategic Software Sales Specialist"
    core.category = "Resume"
    core.comments = (
        "Updated from v13 to add EA Architect / Combined Seller (Aug 2023–Present) "
        "and Strategic Software Sales Specialist (Oct 2019–Present); closed Carolina Select SE at Oct 2019."
    )

    doc.save(OUT)
    print("Wrote", OUT)


if __name__ == "__main__":
    build()
